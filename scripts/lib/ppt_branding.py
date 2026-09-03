"""Generic PowerPoint helpers for ML-Platform diagram exports (no client branding)."""

from __future__ import annotations

import re
import shutil
import subprocess
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

BLANK_LAYOUT = 6
SLIDE_W = 13.333
SLIDE_H = 7.5
FONT = "Calibri"

NAVY = RGBColor(0x20, 0x11, 0x3D)
PURPLE = RGBColor(0x4F, 0x00, 0xCA)
BODY = RGBColor(0x23, 0x1E, 0x33)
SLATE = RGBColor(0x54, 0x65, 0x7A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PAGE_NUM_Y = 7.04


def new_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    return prs


def post_process_pptx(path: Path) -> None:
    """Normalize package for cross-platform PowerPoint compatibility."""
    tmp = path.with_suffix(".tmp.pptx")
    with zipfile.ZipFile(path, "r") as zin:
        with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                name = item.filename
                if "docProps/thumbnail" in name or "printerSettings" in name:
                    continue
                base = name.rsplit("/", 1)[-1]
                if name.startswith("__MACOSX") or base.startswith("._"):
                    continue
                data = zin.read(name)
                if name == "[Content_Types].xml":
                    text = data.decode("utf-8")
                    text = re.sub(r'<Override PartName="/docProps/thumbnail\.jpeg"[^/]*/>\s*', "", text)
                    text = re.sub(
                        r'<Override PartName="/ppt/printerSettings/printerSettings1\.bin"[^/]*/>\s*',
                        "",
                        text,
                    )
                    data = text.encode("utf-8")
                elif name.endswith(".xml") and ("/slides/" in name or "/slideLayouts/" in name):
                    text = data.decode("utf-8")

                    def _clamp_ext(m: re.Match) -> str:
                        cx = max(0, int(m.group(1)))
                        cy = max(0, int(m.group(2)))
                        return f'<a:ext cx="{cx}" cy="{cy}"/>'

                    text = re.sub(r'<a:ext cx="(-?\d+)" cy="(-?\d+)"/>', _clamp_ext, text)
                    data = text.encode("utf-8")
                info = zipfile.ZipInfo(filename=name, date_time=item.date_time)
                info.compress_type = zipfile.ZIP_DEFLATED
                info.create_system = 0
                zout.writestr(info, data)
    shutil.move(tmp, path)
    subprocess.run(["xattr", "-cr", str(path)], check=False)


def _style_run(run, size_pt: float, color=BODY, bold=False) -> None:
    run.font.name = FONT
    run.font.size = Pt(size_pt)
    run.font.color.rgb = color
    run.font.bold = bold


def add_footer(slide, deck_text: str | None = None, slide_num=None) -> None:
    rule = slide.shapes.add_shape(1, Inches(0.36), Inches(6.78), Inches(12.4), Inches(0.01))
    rule.fill.solid()
    rule.fill.fore_color.rgb = SLATE
    rule.line.fill.background()
    if deck_text:
        cap = slide.shapes.add_textbox(Inches(0.36), Inches(6.82), Inches(10.5), Inches(0.18))
        p = cap.text_frame.paragraphs[0]
        p.text = deck_text
        _style_run(p.runs[0] if p.runs else p.add_run(), 8, SLATE)
    if slide_num is not None:
        num = slide.shapes.add_textbox(Inches(12.55), Inches(PAGE_NUM_Y), Inches(0.55), Inches(0.22))
        np = num.text_frame.paragraphs[0]
        np.text = str(slide_num)
        _style_run(np.runs[0], 9, NAVY)
        np.alignment = PP_ALIGN.RIGHT


def add_title_bar(slide, section: str, title: str) -> None:
    del section
    box = slide.shapes.add_textbox(Inches(0.36), Inches(0.28), Inches(12.4), Inches(0.55))
    p = box.text_frame.paragraphs[0]
    p.text = title
    _style_run(p.runs[0], 24, PURPLE, bold=True)
